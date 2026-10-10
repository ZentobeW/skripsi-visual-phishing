"""
Defense Guide DOCX Generator (Natural, Non-Robotic Tone).
Builds a printable, beautifully formatted Microsoft Word version of
docs/PANDUAN_ARGUMENTASI_DOSEN.md with natural, conversational language.
"""

from pathlib import Path
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor

BASE_DIR = Path(__file__).resolve().parent.parent
DOCS_DIR = BASE_DIR / "docs"
OUTPUT_DOCX = DOCS_DIR / "Buku_Panduan_Argumentasi_Sidang_Dosen.docx"


def set_cell_background(cell, fill_hex: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)


def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="12" w:space="0" w:color="000000"/>'
        f'  <w:bottom w:val="single" w:sz="12" w:space="0" w:color="000000"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="D3D3D3"/>'
        f'  <w:insideV w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def format_p(p, before=2, after=6, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=0.0):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line_spacing
    p.alignment = align
    if indent > 0:
        p.paragraph_format.first_line_indent = Inches(indent)


def add_callout(doc, text: str, title: str = "CARA MENJAWAB YANG LUWES & NANCEP"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.0)
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F0F4F8")

    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="1A5276"/>'
        f'  <w:top w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    p_title = cell.paragraphs[0]
    r_t = p_title.add_run(f"🎙️ {title}\n")
    r_t.bold = True
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(10)
    r_t.font.color.rgb = RGBColor(26, 82, 118)

    r_body = p_title.add_run(text)
    r_body.italic = True
    r_body.font.name = "Times New Roman"
    r_body.font.size = Pt(10.5)
    r_body.font.color.rgb = RGBColor(33, 47, 61)
    format_p(p_title, before=4, after=4, align=WD_ALIGN_PARAGRAPH.LEFT)

    p_space = doc.add_paragraph()
    format_p(p_space, after=4)


def add_heading(doc, text: str, level: int):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.bold = True
    if level == 1:
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(26, 82, 118)
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
    elif level == 2:
        run.font.size = Pt(11.5)
        run.font.color.rgb = RGBColor(40, 55, 71)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
    return p


def build_defense_docx():
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    # Title
    p_main = doc.add_paragraph()
    p_main.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_main.add_run(
        "PANDUAN NGOBROL & ARGUMENTASI DENGAN DOSEN\n"
        "POLITEKNIK STATISTIKA STIS"
    )
    r_title.bold = True
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(14)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run(
        "Topik: Deteksi Visual Phishing dengan Siamese Vision Transformer & Adversarial Training"
    )
    r_sub.italic = True
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(10)
    format_p(p_sub, after=18)

    # Section 1
    add_heading(doc, "1. Cara Menjelaskan Inti Skripsi dalam 2 Menit (Elevator Pitch)", level=1)
    pitch_text = (
        '"Jadi begini Pak/Bu, masalah utamanya sebenarnya sederhana tapi dampaknya fatal.\n\n'
        'Masyarakat kita itu rata-rata nggak merhatiin nama domain di browser. Mau link-nya berakhiran .go.id, .xyz, atau .pages.dev, '
        'kebanyakan orang nggak paham bedanya. Tapi ada satu hal yang bikin mereka langsung percaya: begitu mereka buka halaman web '
        'dan melihat ada Logo Garuda, Logo Ditjen Pajak, atau Logo BPS, rasa curiganya langsung hilang. Mereka mikir, "Oh, ini situs resmi '
        'pemerintah," lalu dengan gampang masukin NIK, password, sampai nomor rekening.\n\n'
        'Di sisi lain, cara lama Kominfo memblokir web penipuan lewat sistem blacklist itu punya jeda 1 sampai 3 hari karena nunggu '
        'laporan manual. Padahal web phishing zaman sekarang umurnya cuma 2 sampai 4 jam sebelum dibuang oleh pelaku.\n\n'
        'Nah, satu-satunya hal yang nggak bisa dihindari oleh penyerang adalah: mereka wajib pasang logo instansi biar korbannya percaya.\n\n'
        'Di skripsi ini, saya bikin sistem AI penglihatan (Vision Transformer) yang tugasnya mencocokkan kemiripan logo kementerian secara otomatis. '
        'Bedanya dengan penelitian luar negeri:\n'
        '1. Model saya dilatih tahan banting melawan manipulasi derau tak kasatmata (adversarial attack), jadi nggak gampang dikelabui penyerang.\n'
        '2. Diuji murni pada skenario instansi yang belum pernah dipelajari model sebelumnya (zero-day).\n'
        '3. Dan yang paling penting dari sisi statistika STIS: ambang batas keputusannya diuji ketat pakai kriteria Neyman-Pearson '
        'dan signifikansinya dibuktikan lewat Bootstrap 1.000 kali putaran.\n\n'
        'Output-nya sudah jadi API siap pakai dengan kecepatan respons di bawah 100 milidetik, jadi realistis buat dipakai BSSN atau Komdigi. '
        'Kurang lebih seperti itu intinya, Pak/Bu."'
    )
    add_callout(doc, pitch_text, title="CONTOH PENJELASAN SANTAI & BERBOBOT")

    # Section 2
    add_heading(doc, "2. Kalau Dosen Meragukan: 'Emang Ada Korbannya? Bukannya Phishing Cuma APK di WA?'", level=1)
    q1_text = (
        '"Betul Pak/Bu, sekarang memang banyak yang kirim undangan APK lewat WA. Tapi kalau kita bedah alurnya, pesan WA itu cuma pintunya saja.\n\n'
        'Begitu link-nya diklik, korban tetap dilempar ke halaman web peramban. Nah, di halaman web itulah penipuan visualnya bekerja.\n\n'
        'Apalagi sekarang pelaku nggak perlu sewa server mahal-mahal. Mereka pakai layanan cloud gratisan kayak Cloudflare Pages atau Firebase '
        'yang otomatis dapet HTTPS gembok hijau. Orang awam kalau lihat ada gembok hijau dan ada logo kementerian, mereka langsung ngerasa aman.\n\n'
        'Kasusnya nyata banget di lapangan, Pak/Bu:\n'
        '- Di Pajak, banyak wajib pajak ketipu situs e-Filing palsu berkedok pengembalian lebih bayar.\n'
        '- Di Kemensos, ribuan NIK warga dicuri lewat web tiruan "Cek Bansos" buat dipakai pinjol ilegal.\n'
        '- Di BPS sendiri, sering ada web tiruan rekrutmen mitra statistik yang mungut biaya pendaftaran.\n'
        '- Dan di OJK, logo mereka dicatut di web investasi bodong biar kelihatan resmi.\n\n'
        'Jadi ancamannya bukan cuma orang kehilangan uang, tapi kepercayaan publik ke layanan digital pemerintah yang ikut rusak."'
    )
    add_callout(doc, q1_text, title="CARA MENJAWAB URGENSI MASALAH DENGAN FAKTA")

    # Section 3
    add_heading(doc, "3. Kalau Dosen Nanya Kebaruan: 'Di Luar Negeri Kan Sudah Banyak Riset Begini?'", level=1)
    q2_text = (
        '"Betul sekali Pak/Bu, paper luar seperti Phishpedia di USENIX memang jadi kiblat awal.\n\n'
        'Tapi setelah saya bedah paper mereka, ada satu kelemahan besar: mereka menganggap penyerangnya pasif dan polos. '
        'Model mereka cuma dites pakai logo bersih hasil cropping biasa.\n\n'
        'Padahal di dunia nyata, penyerang yang pintar bisa ngasih sedikit derau piksel halus—namanya serangan PGD. '
        'Di mata manusia logonya tetap kelihatan normal kayak logo BPS, tapi AI luar negeri itu langsung "buta" dan akurasinya '
        'anjlok parah ke belasan persen.\n\n'
        'Nah, celah itulah yang saya selesaikan di skripsi ini:\n'
        '1. Model saya nggak dilatih secara pasif, tapi pakai konsep Min-Max Game Theory. Jadi pas latihan, model dipaksa bertarung '
        'melawan serangan derau PGD biar kebal.\n'
        '2. Arsitekturnya saya naikkan dari CNN biasa ke Vision Transformer (ViT-B/16), jadi AI-nya melihat logo secara menyeluruh, '
        'bukan cuma tekstur lokal.\n'
        '3. Dan dataset luar negeri itu isinya logo minimalis kayak Apple atau Netflix. Belum ada yang meneliti logo instansi Indonesia '
        'yang penuh ornamen lambang heraldik rumit kayak Garuda, padi-kapas, dan lingkaran teks.\n\n'
        'Jadi kita nggak bikin ulang dari nol, tapi memperbaiki kelemahan fatal dari riset dunia yang sudah ada."'
    )
    add_callout(doc, q2_text, title="CARA MENJELASKAN NOVELTY TANPA MENGGURUI")

    # Section 4
    add_heading(doc, "4. Kalau Dosen Nanya Sisi STIS: 'Ini Kayak Skripsi Informatika, Mana Statistikanya?'", level=1)
    q3_text = (
        '"Justru di bagian evaluasinya ini "jiwa STIS"-nya keluar penuh, Pak/Bu.\n\n'
        'Kalau teman-teman di Informatika biasanya puas cuma dengan pamer angka deskriptif satu titik—misal akurasinya 92%, terus selesai. '
        'Nah, di skripsi ini saya perlakukan hasil AI itu sebagai penduga statistik yang punya ketidakpastian.\n\n'
        'Ada tiga hal statistik murni yang saya terapkan:\n'
        '1. Optimisasi Bersyarat Neyman-Pearson: Threshold atau ambang batas vonisnya nggak saya tebak-tebak pakai angka 0.5. '
        'Saya kalibrasi pakai Lema Neyman-Pearson: kita maksimalkan Recall, tapi kita kunci False Positive Rate-nya maksimal 1%. '
        'Kenapa 1%? Karena kalau sistem keamanan sering salah tuduh, analis di BSSN bakal kena "alert fatigue" atau capek ngurusin alarm palsu.\n'
        '2. Inferensi Non-Parametrik Bootstrap 1.000 Kali: Saya nggak mau mengasumsikan datanya berdistribusi normal. Jadi saya lakukan '
        'resampling Bootstrap 1.000 kali buat nyari Selang Kepercayaan 95% dari metrik PR-AUC. Hasilnya terbukti selang kepercayaan '
        'model usulan saya sama sekali nggak tumpang tindih dengan model baseline.\n'
        '3. Uji Hipotesis Komparasi DeLong: Saya lakukan uji hipotesis formal pakai uji DeLong buat ngebuktiin apakah selisih performa model saya '
        'ini beneran beda nyata secara statistik atau cuma kebetulan sampel lagi bagus. Nilai p-value-nya jauh di bawah 0.001, jadi penolakan H0-nya '
        'sangat meyakinkan.\n\n'
        'Jadi metodologi pengujian di skripsi ini murni inferensi statistika, bukan sekadar coba-coba coding."'
    )
    add_callout(doc, q3_text, title="CARA MENUNJUKKAN JIWA KOMPUTASI STATISTIK")

    # Section 5
    add_heading(doc, "5. Kalau Dosen Nanya: 'Kenapa Datanya Harus Brand-Disjoint?'", level=1)
    q4_text = (
        '"Kalau kita acak gambarnya kayak split biasa di machine learning, nanti logo Kemenkeu yang asli ada di data latih, '
        'terus logo Kemenkeu yang tiruan ada di data uji.\n\n'
        'Itu namanya kebocoran data (Data Leakage), Pak/Bu. Modelnya bakal kelihatan pintar banget dapat skor 99%, '
        'padahal aslinya dia cuma menghafal bentuk logo yang sudah pernah dia lihat.\n\n'
        'Di lapangan, serangan yang paling berbahaya itu serangan Zero-Day: tiba-tiba ada web penipu bikin tiruan instansi '
        'yang model AI-nya belum pernah lihat sama sekali pas latihan—misalnya BPS atau OJK.\n\n'
        'Makanya saya pakai metode Brand-Disjoint: instansi yang dipakai buat ujian akhir (BPS, OJK, BSSN, Kejaksaan) bener-bener '
        'diisolasi total. Model nggak boleh ngintip logo itu sepeser pun selama masa pelatihan. Dengan begitu, hasil pengujiannya '
        'jujur dan mencerminkan kondisi lapangan yang sebenarnya."'
    )
    add_callout(doc, q4_text, title="CARA MENJELASKAN BRAND-DISJOINT SECARA INTUITIF")

    # Section 6
    add_heading(doc, "6. Cheatsheet Kata Pengganti Biar Nggak Terdengar Robot", level=1)
    t = doc.add_table(rows=6, cols=2)
    set_table_borders(t)

    t.cell(0, 0).text = "Jangan Pakai Kalimat Kaku Ini (Terlalu Robot) ❌"
    t.cell(0, 0).paragraphs[0].runs[0].bold = True
    set_cell_background(t.cell(0, 0), "FADBD8")

    t.cell(0, 1).text = "Pakai Bahasa Alami Ini Saja (Luwes & Nancep) ✅"
    t.cell(0, 1).paragraphs[0].runs[0].bold = True
    set_cell_background(t.cell(0, 1), "D4EFDF")

    rows_data = [
        ('"Izin menjawab, penelitian ini dilatarbelakangi oleh paradoks kritis..."', '"Jadi begini Pak/Bu, masalah utamanya sebenarnya sederhana tapi dampaknya bahaya..."'),
        ('"Mekanisme pertahanan tradisional mengalami fenomena detection lag..."', '"Cara lama Kominfo itu punya jeda 1 sampai 3 hari karena nunggu laporan manual..."'),
        ('"Model saya mengintegrasikan formulasi optimisasi minimax madry..."', '"Model saya dilatih pakai konsep game theory, jadi pas latihan AI-nya dipaksa bertarung melawan serangan derau biar kebal..."'),
        ('"Kami memitigasi anomali data leakage melalui brand-disjoint partitioning..."', '"Datanya sengaja kita pisah berdasarkan instansi, biar modelnya nggak cuma menghafal logo yang sudah pernah dilihat..."'),
        ('"Untuk membuktikan signifikansi empiris non-parametrik..."', '"Biar nggak cuma klaim angka akurasi doang, saya uji pakai Bootstrap 1.000 kali buat ngebuktiin selisihnya beneran nyata..."'),
    ]

    for i, (kaku, luwes) in enumerate(rows_data):
        t.cell(i + 1, 0).text = kaku
        t.cell(i + 1, 1).text = luwes

    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    try:
        doc.save(str(OUTPUT_DOCX))
        print(f"[SUCCESS] Defense guide DOCX generated: {OUTPUT_DOCX}")
    except PermissionError:
        fallback = OUTPUT_DOCX.parent / "Buku_Panduan_Argumentasi_Sidang_Dosen_Updated.docx"
        doc.save(str(fallback))
        print(f"[SUCCESS] Defense guide DOCX generated (fallback): {fallback}")


if __name__ == "__main__":
    build_defense_docx()
